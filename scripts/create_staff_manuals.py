from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("docs")
BRAND_DARK = "181513"
BRAND_RED = "AD2B25"
CREAM = "FFFDF8"
MUTED = "74685D"


def set_font(run, size, color=BRAND_DARK, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BRAND_DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color in [("Heading 1", 16, BRAND_RED), ("Heading 2", 12.5, BRAND_DARK), ("Heading 3", 11, BRAND_DARK)]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(11)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 3"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 3"].paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("SAKE STREET  |  TABLE ORDERING")
    set_font(header_run, 8.5, MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("INTERNAL STAFF GUIDE  |  Check the live screen before acting")
    set_font(footer_run, 8, MUTED)


def title_block(doc, title, subtitle, revision):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("SAKE STREET")
    set_font(r, 10, BRAND_RED, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, 25, BRAND_DARK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(subtitle)
    set_font(r, 11.5, MUTED)

    band = doc.add_table(rows=1, cols=1)
    band.autofit = False
    cell = band.cell(0, 0)
    set_cell_fill(cell, CREAM)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(revision)
    set_font(r, 9, BRAND_DARK, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, 10.5, BRAND_DARK, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, 10.5)
    else:
        r = p.add_run(text)
        set_font(r, 10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, 10.5)


def add_step(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, 10.5)


def add_warning(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_fill(cell, "F9EAE8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + " ")
    set_font(r, 10, BRAND_RED, bold=True)
    r = p.add_run(text)
    set_font(r, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def detailed_manual():
    doc = Document()
    configure(doc)
    title_block(
        doc,
        "员工操作手册 / Staff Operations SOP",
        "Sake Street QR Table Ordering - front desk, kitchen and manager procedures",
        "Version 1.0  |  15 July 2026  |  Use with sakestreet.aveniq.com.au",
    )

    doc.add_heading("1. 本手册的用途 / Purpose", level=1)
    add_body(doc, "本手册说明员工如何处理顾客扫码下单、厨房订单、菜品售罄、前台结账及日常检查。界面按钮保留英文名称，便于与实际系统一致。")
    add_warning(doc, "原则：", "顾客订单以屏幕上的桌号、项目、数量和备注为准；不要只凭口头信息备餐。")

    doc.add_heading("2. 开店前检查 / Opening checklist", level=1)
    for item in [
        "在一台厨房平板或电脑打开网站，使用 Staff Login 登录厨房或经理账号。",
        "确认页头显示 Open；若显示 Closed，先在 Admin 的 Restaurant profile 改为 Open。",
        "确认 Cloud sync 显示 Connected 或最近同步时间正常；若显示 Offline，先检查店内网络，再通知负责人。",
        "打开 Kitchen 页面，确认声音提醒按需要开启，且屏幕能持续显示订单。",
        "检查 Menu Availability：当天没有的菜必须标记 sold out，避免顾客继续下单。",
        "随机扫描一张桌码，确认能打开对应的桌号和菜单。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("3. 顾客下单流程 / Customer ordering", level=1)
    add_body(doc, "顾客扫描桌上的 QR code 后，系统会显示 Scanned table 与桌号。顾客在菜单点击 +，检查购物车，再点击 Send to Kitchen。")
    for step in [
        "顾客选择菜品。数量可在购物车内用 - / + 调整；备注填写在 Order note。",
        "顾客确认总额和 GST 后点击 Send to Kitchen。",
        "提交成功后购物车清空，顾客可点击 View order status 查看 Received / Preparing / Ready / Served。",
        "厨房屏幕会出现新订单提醒；厨房人员须先核对桌号、数量、备注和过敏原要求。",
    ]:
        add_step(doc, step)
    add_warning(doc, "测试订单：", "必须在备注首行写 TEST ONLY - DO NOT PREPARE，并由经理在厨房屏幕确认后处理或取消。")

    doc.add_heading("4. 厨房处理订单 / Kitchen workflow", level=1)
    add_body(doc, "在 Kitchen 页面，每张订单应按实际出餐进度更新状态。先读备注，再开始备餐。")
    for step in [
        "新订单出现时，确认桌号、所有菜品、数量及 Order note；有过敏原或不清楚的备注时，先向前台确认。",
        "开始制作时，将订单改为 Preparing。",
        "所有项目可出餐时改为 Ready，并通知前台或送餐人员。",
        "送至正确桌位后改为 Served。不要在未送餐前标记 Served。",
        "取消、缺货或需修改订单时，由经理或前台先与顾客确认，再按系统可用操作处理并在备注留下原因。",
    ]:
        add_step(doc, step)
    add_warning(doc, "重要：", "不要用 Clear Local 清除仍需处理的订单。该功能只适用于本机的本地测试/显示清理；正式订单必须先确认已同步并完成处理。")

    doc.add_heading("5. 前台与结账 / Front desk", level=1)
    for step in [
        "在 Front Desk 查找桌号，核对该桌所有未完成订单与总金额。",
        "收款前再次确认是否有未送达、已取消或需要拆单的项目。",
        "按餐厅现行收款方式完成付款；系统中的付款/状态操作必须与实际收款一致。",
        "处理完毕后确认订单不再显示为 open，避免同一桌重复收费。",
    ]:
        add_step(doc, step)

    doc.add_heading("6. 菜品售罄、菜单和照片 / Menu control", level=1)
    for step in [
        "厨房发现某菜不可供应时，立即在 Kitchen 的 Menu Availability 标记 sold out。",
        "恢复供应前，在菜单核实库存或主厨确认后取消 sold out。",
        "价格、名称、过敏原、类别、图片等长期修改只由 Manager / Owner 在 Admin 处理。修改后用顾客页面复查。",
        "不要在营业高峰期大量修改菜单；先由一人负责、另一人复核。",
    ]:
        add_step(doc, step)

    doc.add_heading("7. 常见问题 / Troubleshooting", level=1)
    doc.add_heading("顾客说扫不了 QR code", level=2)
    add_bullet(doc, "请顾客拉近/拉远镜头、提高屏幕亮度或使用相机的扫码功能；必要时提供该桌的 Copy Link。")
    add_bullet(doc, "确认桌码打印清楚、四周保留白边，且对应桌号正确。")
    doc.add_heading("厨房没有收到订单", level=2)
    add_bullet(doc, "先确认顾客是否看到提交成功，以及厨房设备是否在线、登录并停留在 Kitchen 页面。")
    add_bullet(doc, "查看 Cloud sync 状态；若离线，恢复网络后刷新并检查订单。不要让顾客重复提交，先查桌号。")
    doc.add_heading("订单或金额不对", level=2)
    add_bullet(doc, "先暂停出餐/收款，核对桌号、数量、备注与顾客手机。由经理决定修改、取消或补单。")
    doc.add_heading("页面显示旧版本", level=2)
    add_bullet(doc, "刷新浏览器；若为已安装 PWA，完全关闭后重开。仍未恢复时请记录设备、时间和截图。")

    doc.add_heading("8. 收店检查 / Closing checklist", level=1)
    for item in [
        "Kitchen 和 Front Desk 中没有未处理的 New / Preparing / Ready 订单。",
        "确认最后一桌已结账或交接给下一班。",
        "在 Admin 导出 Backup，并保存到餐厅指定位置。备份不包含订单与员工账号。",
        "如明日继续使用，复核餐厅状态、售罄菜品和设备充电。",
        "退出共享设备上的员工账号，关闭不需要的浏览器标签。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("9. 升级或故障上报信息", level=1)
    add_body(doc, "上报时请提供：发生时间、桌号、订单号（如有）、设备类型、网络状态、截图、已经尝试的操作，以及是否影响顾客或厨房。")
    return doc


def quick_guide():
    doc = Document()
    configure(doc)
    title_block(
        doc,
        "员工快速卡 / Quick Staff Card",
        "One-page reference for service and kitchen teams",
        "Sake Street QR Ordering  |  Keep this beside the kitchen screen",
    )

    doc.add_heading("开店 / Start", level=1)
    for item in [
        "Staff Login -> 进入 Kitchen 或 Front Desk。",
        "确认 Open、网络正常、Cloud sync 正常。",
        "Kitchen 打开并开启声音提醒（如需要）。",
        "先把没有库存的菜标记 sold out。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("厨房 / Kitchen", level=1)
    for item in [
        "新单先看：桌号 + 菜品 + 数量 + Order note。",
        "开始做 -> Preparing。",
        "可出餐 -> Ready。",
        "送到正确桌 -> Served。",
        "有过敏原、缺货或不清楚的备注：先停下并问前台。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("前台 / Front Desk", level=1)
    for item in [
        "结账前按桌号核对所有未完成订单。",
        "确认送达/取消/拆单后再收款。",
        "不要因为顾客催促就重复提交订单；先查该桌订单状态。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("三条不能忘 / Non-negotiables", level=1)
    add_warning(doc, "1.", "不要忽略 Order note，尤其是过敏原或去除食材。")
    add_warning(doc, "2.", "不要在正式订单仍未处理时使用 Clear Local。")
    add_warning(doc, "3.", "测试单必须写 TEST ONLY - DO NOT PREPARE。")

    doc.add_heading("有问题时 / If something is wrong", level=1)
    add_body(doc, "先截图并记录桌号、时间和订单状态 -> 检查网络与 Cloud sync -> 通知经理。")
    return doc


def save(doc, filename):
    OUT_DIR.mkdir(exist_ok=True)
    path = OUT_DIR / filename
    doc.save(path)
    print(path.resolve())


if __name__ == "__main__":
    save(detailed_manual(), "Sake-Street-Staff-Operations-SOP.docx")
    save(quick_guide(), "Sake-Street-Staff-Quick-Guide.docx")
